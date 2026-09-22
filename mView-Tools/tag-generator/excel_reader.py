from __future__ import annotations

"""Lectura flexible de Excel para convertir filas a TagRecord."""

from dataclasses import dataclass
from openpyxl import load_workbook
from mview_tag import TagRecord


HEADER_ALIASES = {
    "tag": {"tag", "tag mview", "tag_name", "name", "nombre", "etiqueta"},
    "address": {
        "direccion", "dirección", "direccion plc", "dirección plc",
        "dispositivo plc", "device", "address", "plc", "registro",
    },
    "comment": {
        "comentario", "comment", "descripcion", "descripción", "description",
    },
    "maximum": {"max", "maximo", "máximo", "máx", "maximum"},
    "minimum": {"min", "minimo", "mínimo", "mín", "minimum"},
}


def _norm(value) -> str:
    return " ".join(str(value or "").strip().lower().split())


@dataclass
class SheetInfo:
    name: str
    max_row: int
    max_column: int


def list_sheets(path: str) -> list[SheetInfo]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return [
            SheetInfo(ws.title, ws.max_row or 0, ws.max_column or 0)
            for ws in wb.worksheets
        ]
    finally:
        wb.close()


def read_headers(path: str, sheet_name: str, header_row: int) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        row = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
        return [str(v or "").strip() for v in row]
    finally:
        wb.close()


def detect_header_row(
    path: str,
    sheet_name: str,
    *,
    search_rows: int = 30,
) -> tuple[int, dict[str, int]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        for row_idx, row in enumerate(
            ws.iter_rows(
                min_row=1,
                max_row=min(search_rows, ws.max_row or search_rows),
                values_only=True,
            ),
            start=1,
        ):
            normalized = [_norm(v) for v in row]
            mapping: dict[str, int] = {}
            for field, aliases in HEADER_ALIASES.items():
                for idx, header in enumerate(normalized):
                    if header in aliases:
                        mapping[field] = idx
                        break
            if "tag" in mapping and "address" in mapping:
                return row_idx, mapping
    finally:
        wb.close()

    raise ValueError(
        "No pude detectar una fila de encabezados con columnas TAG y DIRECCIÓN."
    )


def records_from_excel(
    path: str,
    sheet_name: str,
    header_row: int,
    columns: dict[str, int | None],
) -> list[TagRecord]:
    if columns.get("tag") is None or columns.get("address") is None:
        raise ValueError("Las columnas TAG y DIRECCIÓN son obligatorias.")

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        result: list[TagRecord] = []

        def get(row, key: str) -> str:
            idx = columns.get(key)
            if idx is None or idx >= len(row):
                return ""
            value = row[idx]
            if value is None:
                return ""
            if isinstance(value, float) and value.is_integer():
                return str(int(value))
            return str(value).strip()

        for row in ws.iter_rows(
            min_row=header_row + 1,
            max_row=ws.max_row,
            values_only=True,
        ):
            name = get(row, "tag")
            address = get(row, "address")
            if not name and not address:
                continue
            result.append(
                TagRecord(
                    name=name,
                    address=address,
                    comment=get(row, "comment"),
                    maximum=get(row, "maximum"),
                    minimum=get(row, "minimum"),
                )
            )
        return result
    finally:
        wb.close()
